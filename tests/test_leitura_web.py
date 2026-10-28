"""Leitura correta de HTML e derivados, e da codificação de todo arquivo de texto.

Dois defeitos que esta suíte prende. Primeiro: o HTML entrava no índice
como marcação crua — `<head>`, CSS e JavaScript ocupavam os caracteres por
página, e o texto tachado (dispositivo revogado, na legislação do Planalto)
ficava indistinguível do vigente. Segundo: todo arquivo de texto era lido
como UTF-8 com `errors="replace"`, e os acervos brasileiros em Windows-1252
viravam "Presid�ncia da Rep�blica".
"""

from __future__ import annotations

from pathlib import Path

import pytest

from gclaude_indexer.config import load_config
from gclaude_indexer.conversion import convert
from gclaude_indexer.db import connect, init_schema
from gclaude_indexer.file_types import category_of_extension
from gclaude_indexer.html_text import html_to_text, without_struck_text
from gclaude_indexer.scanning import scan
from gclaude_indexer.text_decoding import decode_bytes, decode_markup

FIXTURES = Path(__file__).parent / "fixtures" / "web"


# --- decodificação ----------------------------------------------------------


def test_utf8_sem_declaracao_e_lido_como_utf8():
    texto, codificacao = decode_bytes("Presidência da República".encode("utf-8"))
    assert texto == "Presidência da República"
    assert codificacao == "utf-8"


def test_bom_utf8_e_removido():
    texto, codificacao = decode_bytes(b"\xef\xbb\xbf" + "ação".encode("utf-8"))
    assert texto == "ação"
    assert codificacao == "utf-8-sig"


def test_bom_utf16_decide_a_codificacao():
    texto, _ = decode_bytes("Código Civil".encode("utf-16"))
    assert texto == "Código Civil"


def test_windows_1252_sem_declaracao_nao_vira_losango():
    bruto = "Presidência da República — “Lei” nº 10.147".encode("cp1252")
    texto, codificacao = decode_bytes(bruto)
    assert texto == "Presidência da República — “Lei” nº 10.147"
    assert "�" not in texto
    assert codificacao == "cp1252"


def test_iso_8859_1_declarado_e_lido_como_windows_1252():
    # Regra do WHATWG: o rótulo "iso-8859-1" quer dizer windows-1252 — as
    # aspas curvas 0x93/0x94 do Word só existem nesta.
    bruto = "“vigência”".encode("cp1252")
    texto, codificacao = decode_bytes(bruto, declared="ISO-8859-1")
    assert texto == "“vigência”"
    assert codificacao == "cp1252"


def test_utf8_declarado_mas_falso_cai_na_deteccao():
    bruto = "Subchefia para Assuntos Jurídicos".encode("cp1252")
    texto, _ = decode_bytes(bruto, declared="utf-8")
    assert texto == "Subchefia para Assuntos Jurídicos"


def test_rotulo_desconhecido_nao_derruba_a_leitura():
    texto, _ = decode_bytes("ação".encode("utf-8"), declared="x-inventado-123")
    assert texto == "ação"


def test_declaracao_de_outra_codificacao_e_respeitada():
    bruto = "Статья 1".encode("cp1251")
    texto, codificacao = decode_bytes(bruto, declared="windows-1251")
    assert texto == "Статья 1"
    assert codificacao == "cp1251"


def test_bytes_vazios_dao_texto_vazio():
    assert decode_bytes(b"") == ("", "utf-8")


def test_lixo_binario_nunca_levanta_excecao():
    texto, _ = decode_bytes(bytes(range(256)) * 4)
    assert isinstance(texto, str)


def test_meta_charset_do_html_e_lido():
    bruto = '<html><head><meta charset="windows-1251"></head><body>Закон</body></html>'.encode("cp1251")
    texto, codificacao = decode_markup(bruto)
    assert "Закон" in texto
    assert codificacao == "cp1251"


def test_meta_http_equiv_do_html_e_lido():
    bruto = (
        '<meta http-equiv="Content-Type" content="text/html; charset=iso-8859-1">'
        "<p>Presidência</p>"
    ).encode("cp1252")
    texto, _ = decode_markup(bruto)
    assert "Presidência" in texto


def test_declaracao_xml_e_lida():
    bruto = '<?xml version="1.0" encoding="ISO-8859-1"?><lei>Ementa: dispõe</lei>'.encode("latin-1")
    texto, _ = decode_markup(bruto)
    assert "dispõe" in texto


# --- HTML para texto ---------------------------------------------------------


def test_cabecalho_script_e_estilo_somem():
    texto = html_to_text(
        "<html><head><title>L1</title><style>p{color:red}</style>"
        "<script>var x = 1;</script></head><body><p>Art. 1º Vigente.</p></body></html>"
    )
    assert texto == "Art. 1º Vigente."


def test_tag_strike_vira_marcador_de_tachado():
    assert html_to_text("<p><strike>Art. 2º Revogado.</strike></p>") == "~~Art. 2º Revogado.~~"


def test_tag_s_e_del_tambem_sao_tachado():
    assert html_to_text("<p><s>a</s> e <del>b</del></p>") == "~~a~~ e ~~b~~"


def test_estilo_inline_line_through_e_tachado():
    html = (
        '<p><span style="font-weight: normal; text-decoration:line-through">'
        "§ 1º A eleição de foro</span></p>"
    )
    assert html_to_text(html) == "~~§ 1º A eleição de foro~~"


def test_classe_css_com_line_through_e_tachado():
    html = (
        "<style>.revogado { text-decoration: line-through }</style>"
        '<p class="Artigo revogado">Art. 3º Texto antigo.</p><p>Art. 4º Novo.</p>'
    )
    assert html_to_text(html) == "~~Art. 3º Texto antigo.~~\nArt. 4º Novo."


def test_tachado_em_varios_paragrafos_marca_cada_linha():
    html = "<strike><p>Art. 5º Um.</p><p>Art. 6º Dois.</p></strike><p>Art. 7º Três.</p>"
    assert html_to_text(html) == "~~Art. 5º Um.~~\n~~Art. 6º Dois.~~\nArt. 7º Três."


def test_tachado_parcial_na_mesma_linha():
    html = "<p>Art. 8º <strike>Fica revogado</strike> Fica mantido o inciso.</p>"
    assert html_to_text(html) == "Art. 8º ~~Fica revogado~~ Fica mantido o inciso."


def test_marcacao_mal_aninhada_nao_vaza_o_tachado():
    html = "<p><strike><font face=Arial>revogado</strike></font> vigente</p><p>outro</p>"
    assert html_to_text(html) == "~~revogado~~ vigente\noutro"


def test_entidades_e_espaco_inquebravel_viram_texto_normal():
    assert html_to_text("<p>Art.&nbsp;1&ordm;&nbsp;&nbsp;A&nbsp;&amp;&nbsp;B&#8212;C</p>") == "Art. 1º A & B—C"


def test_sobrescrito_o_e_a_viram_ordinais():
    assert html_to_text("<p>LEI N<sup>o</sup> 10.147 e 1<sup><u>a</u></sup> vez</p>") == "LEI Nº 10.147 e 1ª vez"


def test_quebras_de_linha_e_titulos_e_listas():
    html = "<h1>Título</h1><p>linha 1<br>linha 2</p><ul><li>item a</li><li>item b</li></ul>"
    assert html_to_text(html) == "# Título\nlinha 1\nlinha 2\n- item a\n- item b"


def test_celulas_de_tabela_ficam_separadas():
    html = "<table><tr><td>Lei</td><td>Data</td></tr><tr><td>10.147</td><td>2000</td></tr></table>"
    assert html_to_text(html) == "Lei | Data\n10.147 | 2000"


def test_pre_preserva_os_espacos():
    assert html_to_text("<pre>a   b\n  c</pre>") == "a   b\n  c"


def test_espacos_e_linhas_em_branco_repetidas_sao_compactados():
    assert html_to_text("<p>  a \n\n   b  </p>\n\n\n<p></p><p>c</p>") == "a b\nc"


def test_html_quebrado_nunca_levanta_excecao():
    texto = html_to_text("<p><strike><b>sem fechar <<< &&& <!-- comentário sem fim")
    assert "sem fechar" in texto


def test_texto_sem_tachado_remove_so_os_trechos_marcados():
    texto = "Art. 1º Vigente.\n~~Art. 2º Revogado.~~\nArt. 3º ~~velho~~ novo."
    assert without_struck_text(texto) == "Art. 1º Vigente.\n\nArt. 3º  novo."


def test_trecho_real_do_planalto_sai_limpo_e_com_tachado():
    bruto = (FIXTURES / "planalto_l10147_trecho.htm").read_bytes()
    texto, _ = decode_markup(bruto)
    limpo = html_to_text(texto)

    assert "Presidência" in limpo
    assert "República" in limpo
    assert "LEI Nº 10.147" in limpo
    assert "~~Art. 1º" in limpo
    assert "�" not in limpo
    assert "<" not in limpo
    assert "FrontPage" not in limpo


# --- tipos de arquivo --------------------------------------------------------


@pytest.mark.parametrize("extensao", [".htm", ".html", ".xhtml", ".xht", ".shtml", ".mht", ".mhtml"])
def test_derivados_de_html_sao_conteudo_web(extensao):
    assert category_of_extension(extensao) == "web_dados"


# --- conversão de ponta a ponta ---------------------------------------------


def _projeto(tmp_path, extensoes):
    origem = tmp_path / "origem"
    origem.mkdir()
    saida = tmp_path / "saida"
    config = load_config({
        "name": "Leitura web",
        "source_folder": str(origem),
        "output_folder": str(saida),
        "extensions": extensoes,
    })
    conn = connect(saida / "project.db")
    init_schema(conn)
    return origem, saida, config, conn


def _converter(conn, config, saida, nome_txt):
    scan(conn, config)
    resultado = convert(conn, config)
    assert resultado.failed == 0
    return (saida / "converted" / nome_txt).read_text(encoding="utf-8")


def test_htm_do_planalto_e_convertido_em_texto_limpo(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["web_dados"])
    (origem / "l10147.htm").write_bytes((FIXTURES / "planalto_l10147_trecho.htm").read_bytes())

    texto = _converter(conn, config, saida, "l10147.txt")

    assert "Presidência da República" in texto
    assert "~~Art. 1º" in texto
    assert "<" not in texto
    conn.close()


def test_txt_em_windows_1252_e_lido_corretamente(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["text"])
    (origem / "nota.txt").write_bytes("Ementa: dispõe sobre a contribuição".encode("cp1252"))

    assert _converter(conn, config, saida, "nota.txt") == "Ementa: dispõe sobre a contribuição"
    conn.close()


def test_xml_que_nao_e_html_fica_literal_mas_decodificado(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["web_dados"])
    (origem / "dados.xml").write_bytes(
        '<?xml version="1.0" encoding="ISO-8859-1"?><norma><ementa>Dispõe</ementa></norma>'.encode("latin-1")
    )

    texto = _converter(conn, config, saida, "dados.txt")
    assert "<ementa>Dispõe</ementa>" in texto
    conn.close()


def test_xhtml_passa_pelo_conversor_de_html(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["web_dados"])
    (origem / "pagina.xhtml").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml"><body><p><del>velho</del> novo</p></body></html>',
        encoding="utf-8",
    )

    assert _converter(conn, config, saida, "pagina.txt") == "~~velho~~ novo"
    conn.close()


def test_mhtml_salvo_pelo_navegador_e_lido(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["web_dados"])
    html = '<html><body><p>Presidência</p><p><strike>Art. 9º</strike></p></body></html>'
    corpo = html.encode("cp1252")
    import quopri

    (origem / "pagina.mht").write_bytes(
        b"MIME-Version: 1.0\r\n"
        b'Content-Type: multipart/related; boundary="XYZ"\r\n\r\n'
        b"--XYZ\r\n"
        b'Content-Type: text/html; charset="windows-1252"\r\n'
        b"Content-Transfer-Encoding: quoted-printable\r\n\r\n"
        + quopri.encodestring(corpo)
        + b"\r\n--XYZ--\r\n"
    )

    assert _converter(conn, config, saida, "pagina.txt") == "Presidência\n~~Art. 9º~~"
    conn.close()


def test_email_so_com_html_em_latin1_e_lido(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["email"])
    (origem / "aviso.eml").write_bytes(
        b"From: a@b.c\r\nSubject: Aviso\r\nMIME-Version: 1.0\r\n"
        b'Content-Type: text/html; charset="iso-8859-1"\r\n\r\n'
        + "<p>Revogação: <s>art. 2º</s></p>".encode("cp1252")
    )

    texto = _converter(conn, config, saida, "aviso.txt")
    assert "Revogação: ~~art. 2º~~" in texto
    assert "<p>" not in texto
    conn.close()


def test_email_com_charset_desconhecido_nao_falha(tmp_path):
    origem, saida, config, conn = _projeto(tmp_path, ["email"])
    (origem / "estranho.eml").write_bytes(
        b"From: a@b.c\r\nSubject: X\r\nMIME-Version: 1.0\r\n"
        b'Content-Type: text/plain; charset="x-nao-existe"\r\n\r\n'
        + "ação".encode("utf-8")
    )

    assert "ação" in _converter(conn, config, saida, "estranho.txt")
    conn.close()


def test_docx_com_tachado_marca_o_trecho(tmp_path):
    from docx import Document

    origem, saida, config, conn = _projeto(tmp_path, ["docx"])
    documento = Document()
    paragrafo = documento.add_paragraph("Art. 1º ")
    paragrafo.add_run("Revogado.").font.strike = True
    paragrafo.add_run(" Vigente.")
    documento.add_paragraph().add_run("Art. 2º Todo revogado.").font.double_strike = True
    documento.save(origem / "lei.docx")

    texto = _converter(conn, config, saida, "lei.txt")
    assert "Art. 1º ~~Revogado.~~ Vigente." in texto
    assert "~~Art. 2º Todo revogado.~~" in texto
    conn.close()


def test_docx_nao_perde_o_texto_de_hiperlinks(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    origem, saida, config, conn = _projeto(tmp_path, ["docx"])
    documento = Document()
    paragrafo = documento.add_paragraph("Ver ")
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    texto_link = OxmlElement("w:t")
    texto_link.text = "Lei 8.112"
    run.append(texto_link)
    link.append(run)
    paragrafo._p.append(link)
    assert link.tag == qn("w:hyperlink")
    documento.save(origem / "ref.docx")

    assert "Ver Lei 8.112" in _converter(conn, config, saida, "ref.txt")
    conn.close()


def test_pptx_com_tachado_marca_o_trecho(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    origem, saida, config, conn = _projeto(tmp_path, ["pptx"])
    apresentacao = Presentation()
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[6])
    caixa = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragrafo = caixa.text_frame.paragraphs[0]
    paragrafo.add_run().text = "Norma "
    riscado = paragrafo.add_run()
    riscado.text = "antiga"
    riscado.font._rPr.set("strike", "sngStrike")
    apresentacao.save(origem / "aula.pptx")

    assert "Norma ~~antiga~~" in _converter(conn, config, saida, "aula.txt")
    conn.close()


def test_xlsx_com_celula_tachada_marca_o_valor(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    origem, saida, config, conn = _projeto(tmp_path, ["xlsx"])
    planilha = Workbook()
    folha = planilha.active
    folha["A1"] = "Lei 1"
    folha["B1"] = "revogada"
    folha["B1"].font = Font(strike=True)
    planilha.save(origem / "tabela.xlsx")

    assert "Lei 1\t~~revogada~~" in _converter(conn, config, saida, "tabela.txt")
    conn.close()


# --- quem lê o marcador sabe o que ele quer dizer ---------------------------


def test_motor_de_regras_ignora_data_tachada():
    from gclaude_indexer.classification import WindowPage
    from gclaude_indexer.engine_rules import RulesEngine, load_rules

    motor = RulesEngine(load_rules())
    pagina = WindowPage(
        reference="p. 1", file_name="lei.htm", has_table=False, image_count=0,
        text="~~LEI Nº 1, DE 3 DE MARÇO DE 1990~~\nLEI Nº 2, DE 5 DE MAIO DE 2001\nDispõe sobre algo.",
    )
    pecas = motor.classify([pagina])
    assert pecas[0].date == "2001-05-05"


def test_prompt_local_explica_o_tachado_so_quando_ha_tachado():
    from gclaude_indexer.classification import WindowPage
    from gclaude_indexer.engine_local import _build_page_prompt, _build_prompt

    com = [WindowPage(reference="p. 1", file_name="a.htm", text="~~Art. 1º~~", has_table=False, image_count=0)]
    sem = [WindowPage(reference="p. 1", file_name="a.pdf", text="Art. 1º", has_table=False, image_count=0)]

    for construir in (_build_prompt, _build_page_prompt):
        assert "tachado" in construir(com)
        assert "tachado" not in construir(sem)


@pytest.mark.parametrize("idioma, palavra", [("pt", "tachado"), ("en", "struck"), ("es", "tachado")])
def test_claude_md_explica_o_tachado_nos_tres_idiomas(tmp_path, idioma, palavra):
    from gclaude_indexer.windows_prep import generate_claude_md

    config = load_config({
        "name": "X", "source_folder": str(tmp_path), "output_folder": str(tmp_path / "saida"),
    })
    (tmp_path / "saida").mkdir()
    texto = generate_claude_md(config, idioma).read_text(encoding="utf-8")
    assert "~~" in texto
    assert palavra in texto.lower()


@pytest.mark.parametrize("idioma", ["pt", "en", "es"])
def test_instrucoes_do_projeto_explicam_o_tachado_quando_o_acervo_tem(tmp_path, idioma):
    from gclaude_indexer.artifacts import generate_project_instructions_md
    from gclaude_indexer.i18n import translate

    origem, saida, config, conn = _projeto(tmp_path, ["web_dados"])
    caminho = generate_project_instructions_md(conn, config, idioma)
    regra = translate(idioma, "artifact.project_instructions.struck_text_rule")
    assert regra not in caminho.read_text(encoding="utf-8")

    conn.execute(
        "INSERT INTO file (relative_path, name, extension, size, mtime, sha256, status) "
        "VALUES ('a.htm', 'a.htm', 'htm', 1, 0, 'x', 'extracted')"
    )
    file_id = conn.execute("SELECT id FROM file").fetchone()[0]
    conn.execute(
        "INSERT INTO page (file_id, number, reference, char_count, image_count, has_table, text) "
        "VALUES (?, 1, 'p. 1', 10, 0, 0, '~~Art. 1º~~')",
        (file_id,),
    )
    conn.commit()
    caminho = generate_project_instructions_md(conn, config, idioma)
    assert regra in caminho.read_text(encoding="utf-8")
    conn.close()
